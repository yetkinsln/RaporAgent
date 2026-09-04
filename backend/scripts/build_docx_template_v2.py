"""Eski kurumsal DOC'tan dönüştürülen boş DOCX'i güvenli Jinja şablonuna çevirir.

Bu araç doldurulmuş örnek belgeyi okumaz. Yalnızca boş şablonun paragraf,
numaralandırma ve sayfa biçimini koruyup açık yer tutucular ekler.
"""

from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Inches


def _replace_runs(paragraph, pieces: list[tuple[str, bool | None, bool | None]]) -> None:
    """Paragraf özelliklerini korurken metin koşularını yeniden kurar."""

    source_rpr = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    for text, bold, underline in pieces:
        run = paragraph.add_run(text)
        if source_rpr is not None:
            run._r.insert(0, deepcopy(source_rpr))
        if bold is not None:
            run.bold = bold
        if underline is not None:
            run.underline = underline


def build_template(source: Path, destination: Path) -> None:
    document = Document(source)
    paragraphs = document.paragraphs
    if len(paragraphs) != 21 or len(document.tables) != 0 or len(document.sections) != 1:
        raise ValueError("Boş kurumsal şablonun beklenen 21 paragraflık tek sayfa yapısı değişmiş.")

    _replace_runs(
        paragraphs[0],
        [
            ("SAYI: ", True, None),
            ("{{ document.number }}", False, None),
            ("\t", False, None),
            ("{{ document.date }}", False, None),
        ],
    )
    paragraphs[0].paragraph_format.tab_stops.add_tab_stop(Inches(7.2), WD_TAB_ALIGNMENT.RIGHT)
    _replace_runs(
        paragraphs[1],
        [("KONU:", True, None), (" {{ patient.full_name }}", False, None)],
    )
    _replace_runs(
        paragraphs[3],
        [
            ("İLGİ:", True, None),
            ("\t", False, None),
            ("a)", True, None),
            (" {{ request.sender_authority }}’nin {{ request.letter_date }} tarih ve {{ request.letter_number }} nolu yazısı.", False, None),
        ],
    )
    _replace_runs(
        paragraphs[4],
        [
            ("b)", True, None),
            (" {{ medical.issuing_institution }}’nin {{ medical.report_date }} tarih ve {{ medical.report_number }} sayılı genel adli muayene raporu.", False, None),
        ],
    )
    _replace_runs(
        paragraphs[5],
        [
            (
                "İlgi (a) yazı ile başvuran {{ patient.full_name }} hakkında {{ request.incident }} nedeniyle {{ request.requested_questions }}.",
                False,
                None,
            )
        ],
    )
    _replace_runs(
        paragraphs[6],
        [
            (
                "{{ patient.full_name }} adına düzenlenen ilgi (b) formda; {{ medical.incident_type }} nedeniyle başvurduğu, "
                "{{ medical.clinical_findings }}; BTM ile {{ medical.btm_record_phrase }} ve hayati tehlikesinin "
                "{{ medical.life_threat_record_phrase }} kayıtlıdır.",
                False,
                None,
            )
        ],
    )
    _replace_runs(
        paragraphs[7],
        [
            (
                "{{ patient.full_name }} adına {{ medical.issuing_institution }} tarafından düzenlenen {{ medical.report_date }} "
                "tarih ve {{ medical.protocol_number }} protokol nolu kayıtta; {{ medical.general_forensic_exam_summary }} kayıtlıdır.",
                False,
                None,
            )
        ],
    )
    _replace_runs(
        paragraphs[8],
        [
            (
                "{{ patient.full_name }}’ın {{ manual.examination_date }} tarihinde saat {{ manual.examination_time }}’te Adli Tıp "
                "Anabilim Dalımızda yapılan muayenesinde; öyküsünde; {{ manual.history }} ifade ettiği; halen yapılan "
                "muayenesinde {{ manual.current_exam_findings }} tespit edilmiştir.",
                False,
                None,
            )
        ],
    )
    _replace_runs(
        paragraphs[9],
        [
            ("SONUÇ:", True, True),
            (
                " İlgili olay nedeniyle Adli Tıp Anabilim Dalımıza gönderilen, {{ patient.tckn }} TC nolu, "
                "{{ patient.birth_info }} doğumlu {{ patient.full_name }}’ın genel adli muayene raporu ve tıbbi evraklarının "
                "tetkiki ile muayenesinden elde edilerek yukarıya kaydedilen bilgi ve bulgulara göre, şahısta mevcut "
                "{{ medical.clinical_findings }} ile ilişkili arızasının;",
                False,
                False,
            ),
        ],
    )
    _replace_runs(
        paragraphs[10],
        [("Şahsın yaşamını ", False, None), ("TEHLİKEYE {{ medical.life_threat_conclusion }},", True, None)],
    )
    _replace_runs(
        paragraphs[11],
        [
            ("Basit tıbbi müdahale ile ", False, None),
            ("giderilebilecek derecede hafif nitelikte {{ medical.btm_conclusion }} ", True, None),
            ("kanaatlerimizi bildirir rapordur.", False, None),
        ],
    )
    _replace_runs(paragraphs[19], [("EKİ:", True, True)])
    _replace_runs(paragraphs[20], [("{{ patient.full_name }}’e ait tıbbi evrakı", False, None)])

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()
    build_template(args.source, args.destination)


if __name__ == "__main__":
    main()
